from io import BytesIO

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from app.docs.dotted import dotted_letter_png
from app.docs.common import (
    HEART,
    GUIDE,
    INK,
    MUTED,
    add_handwriting_lines,
    apply_letter_page,
    center,
    set_cell_margins,
    set_row_height,
    shade_cell,
    style_run,
)
from app.lesson.daily import counting_frames
from app.models import DailyLesson, Word


def build_worksheet(lesson: DailyLesson) -> bytes:
    document = Document()
    apply_letter_page(document.sections[0])
    _student_page(document, lesson)
    document.add_page_break()
    _key_page(document, lesson)
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _student_page(document: Document, lesson: DailyLesson) -> None:
    letter = lesson.letter
    title = document.add_paragraph()
    style_run(title.add_run(letter.letter.upper()), size=48, color=INK, bold=True)
    style_run(title.add_run(f"  {letter.cue}"), size=16, color=MUTED)
    meta = document.add_paragraph()
    style_run(meta.add_run(f"{lesson.date.isoformat()}    Name ____________________"), size=12, color=INK)

    _writing_block(document, letter.letter.upper(), "Write the uppercase letter.")
    _writing_block(document, letter.letter, "Write the lowercase letter.")

    heading = document.add_paragraph()
    style_run(heading.add_run("Sound it out"), size=16, color=INK, bold=True)
    for word in lesson.heart_words + lesson.sight_words:
        _sound_line(document, word)

    bank = document.add_paragraph()
    names = ", ".join(word.display for word in lesson.heart_words + lesson.sight_words)
    style_run(bank.add_run(f"Find these words: {names}"), size=14, color=INK, bold=True)
    _grid(document, lesson, shaded=False)
    document.add_page_break()
    math_title = document.add_paragraph()
    style_run(math_title.add_run("Math sheet"), size=28, color=INK, bold=True)
    math_meta = document.add_paragraph()
    style_run(
        math_meta.add_run(f"{lesson.date.isoformat()}    Name ____________________"),
        size=12,
        color=INK,
    )
    _writing_block(document, str(lesson.number), "Write the number.")
    _counting_block(document, lesson)


def _key_page(document: Document, lesson: DailyLesson) -> None:
    heading = document.add_paragraph()
    style_run(heading.add_run("Grown-up key"), size=20, color=INK, bold=True)
    _grid(document, lesson, shaded=True)
    by_id = {word.id: word for word in lesson.heart_words + lesson.sight_words}
    for place in lesson.hunt.placements:
        line = document.add_paragraph()
        word = by_id[place.word_id]
        style_run(
            line.add_run(
                f"{word.display}: row {place.row + 1}, column {place.col + 1}, {place.direction}"
            ),
            size=14,
            color=INK,
        )


def _writing_block(document: Document, glyph: str, prompt: str) -> None:
    prompt_paragraph = document.add_paragraph()
    style_run(prompt_paragraph.add_run(prompt), size=14, color=INK, bold=True)
    table = document.add_table(rows=1, cols=5)
    table.autofit = False
    widths = (Inches(1.35), Inches(1.35), Inches(1.35), Inches(1.55), Inches(1.55))
    _letter_cell(table.cell(0, 0), widths[0], glyph, highlighted=False)
    _letter_cell(table.cell(0, 1), widths[1], glyph, highlighted=True)
    _letter_cell(table.cell(0, 2), widths[2], glyph, highlighted=True)
    dotted = dotted_letter_png(glyph)
    for index in range(2):
        _blank_cell(table.cell(0, 3 + index), widths[3 + index], dotted)


def _letter_cell(cell, width, glyph: str, highlighted: bool) -> None:
    cell.width = width
    cell.vertical_alignment = WD_ALIGN_VERTICAL.BOTTOM
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = paragraph.add_run(glyph)
    style_run(run, size=40, color=GUIDE if highlighted else INK, bold=True)
    add_handwriting_lines(paragraph)


def _blank_cell(cell, width, dotted: bytes) -> None:
    cell.width = width
    cell.vertical_alignment = WD_ALIGN_VERTICAL.BOTTOM
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.add_run().add_picture(BytesIO(dotted), width=Inches(0.62))
    add_handwriting_lines(paragraph)


def _counting_block(document: Document, lesson: DailyLesson) -> None:
    heading = document.add_paragraph()
    style_run(heading.add_run("Count the dots. How many?"), size=14, color=INK, bold=True)
    for frame in counting_frames(lesson.count):
        table = document.add_table(rows=2, cols=5)
        table.autofit = False
        for index, filled in enumerate(frame):
            cell = table.cell(index // 5, index % 5)
            cell.width = Inches(0.38)
            paragraph = cell.paragraphs[0]
            center(paragraph)
            style_run(paragraph.add_run("●" if filled else "○"), size=16, color=INK)
    prompt = document.add_paragraph()
    style_run(prompt.add_run("Write how many."), size=14, color=INK, bold=True)
    _empty_guides(document, 1)
    follow = document.add_paragraph()
    style_run(follow.add_run("Count on. Write the next numbers."), size=14, color=INK, bold=True)
    table = document.add_table(rows=1, cols=4)
    table.autofit = False
    first = table.cell(0, 0)
    first.width = Inches(1.2)
    paragraph = first.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    style_run(paragraph.add_run(str(lesson.number)), size=28, color=INK, bold=True)
    add_handwriting_lines(paragraph)
    for index in range(1, 4):
        cell = table.cell(0, index)
        cell.width = Inches(1.2)
        line = cell.paragraphs[0]
        add_handwriting_lines(line)
    _skip_row(document, lesson)


def _skip_row(document: Document, lesson: DailyLesson) -> None:
    heading = document.add_paragraph()
    style_run(heading.add_run(f"Skip count by {lesson.skip_by}s."), size=14, color=INK, bold=True)
    table = document.add_table(rows=1, cols=len(lesson.skip_shown) + 3)
    table.autofit = False
    for index, number in enumerate(lesson.skip_shown):
        cell = table.cell(0, index)
        cell.width = Inches(0.9)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        style_run(paragraph.add_run(str(number)), size=22, color=INK, bold=True)
        add_handwriting_lines(paragraph)
    for index in range(3):
        cell = table.cell(0, len(lesson.skip_shown) + index)
        cell.width = Inches(0.9)
        add_handwriting_lines(cell.paragraphs[0])


def _empty_guides(document: Document, count: int) -> None:
    table = document.add_table(rows=1, cols=count)
    table.autofit = False
    for index in range(count):
        cell = table.cell(0, index)
        cell.width = Inches(1.3)
        add_handwriting_lines(cell.paragraphs[0])


def _sound_line(document: Document, word: Word) -> None:
    paragraph = document.add_paragraph()
    for index, piece in enumerate(word.graphemes):
        if index:
            style_run(paragraph.add_run("  "), size=16, color=INK)
        label = piece.letters if piece.phoneme == "-" else f"{piece.letters} {piece.phoneme}"
        style_run(
            paragraph.add_run(label),
            size=16,
            color=HEART if piece.heart else INK,
            bold=True,
            underline=piece.heart,
        )
    sentence = document.add_paragraph()
    style_run(sentence.add_run(word.sentence), size=12, color=MUTED, italic=True)


def _grid(document: Document, lesson: DailyLesson, shaded: bool) -> None:
    rows = lesson.hunt.rows
    table = document.add_table(rows=len(rows), cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    highlights = _highlights(lesson) if shaded else set()
    for r, text in enumerate(rows):
        set_row_height(table.rows[r], 360, "atLeast")
        for c, char in enumerate(text):
            cell = table.rows[r].cells[c]
            cell.width = Inches(0.7)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell, top=20, bottom=20, left=20, right=20)
            paragraph = cell.paragraphs[0]
            center(paragraph)
            style_run(paragraph.add_run(char), size=14, color=INK, bold=True)
            if (r, c) in highlights:
                shade_cell(cell, "F4E2A0")


def _highlights(lesson: DailyLesson) -> set[tuple[int, int]]:
    by_id = {word.id: word for word in lesson.heart_words + lesson.sight_words}
    cells: set[tuple[int, int]] = set()
    for place in lesson.hunt.placements:
        text = by_id[place.word_id].display.lower()
        for offset in range(len(text)):
            row = place.row + (offset if place.direction == "down" else 0)
            col = place.col + (offset if place.direction == "across" else 0)
            cells.add((row, col))
    return cells
