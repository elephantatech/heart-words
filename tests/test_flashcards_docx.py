from io import BytesIO

from docx import Document
from docx.oxml.ns import qn

from app.docs.common import CARD_ROW_TWIPS, HEART
from app.docs.flashcards import build_flashcards
from app.words.load import load_catalog


def _open(words):
    return Document(BytesIO(build_flashcards(words)))


def test_card_rows_fit_a_letter_page():
    document = _open([load_catalog().get("said")])
    height = document.tables[0].rows[0]._tr.trPr.find(qn("w:trHeight"))
    assert height.get(qn("w:val")) == str(CARD_ROW_TWIPS)
    assert height.get(qn("w:hRule")) == "exact"
    margins = document.tables[0].rows[0].cells[0]._tc.tcPr.find(qn("w:tcMar"))
    assert margins.find(qn("w:top")).get(qn("w:w")) == "0"
    assert margins.find(qn("w:bottom")).get(qn("w:w")) == "0"


def test_said_back_keeps_quotes_and_marks_the_heart():
    document = _open([load_catalog().get("said")])
    front_runs = document.tables[0].rows[0].cells[0].paragraphs[0].runs
    assert [run.text for run in front_runs] == ["s", "ai", "d"]
    assert front_runs[1].font.color.rgb == HEART
    back = document.tables[1].rows[0].cells[1]
    assert 'Dad said, "Time for bed."' in back.text


def test_i_card_prints_the_capital_letter():
    document = _open([load_catalog().get("i")])
    assert document.tables[0].rows[0].cells[0].text.startswith("I")
